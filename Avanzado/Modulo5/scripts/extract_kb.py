# -*- coding: utf-8 -*-
"""Extrae el texto de los .docx de /documentacion y lo deja en kb_b64.txt
(un item por documento, {source, file, text}) para que lo use la ingesta.

Requiere: pip install python-docx
"""
import os, glob, json, base64, sys
try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass
from docx import Document
from docx.oxml.ns import qn
import _config as C

EXCLUDE = {"00_Indice.docx", "Documentacion_Completa_LoDeTincho.docx"}

def doc_to_text(path):
    d = Document(path); out = []
    para_map = {p._p: p for p in d.paragraphs}
    tbl_map = {t._tbl: t for t in d.tables}
    for child in d.element.body.iterchildren():
        if child.tag == qn('w:p'):
            p = para_map.get(child)
            if p is None or not p.text.strip():
                continue
            st = (p.style.name or "") if p.style else ""
            t = p.text.strip()
            out.append("# " + t if st == "Title" else
                       "## " + t if st == "Heading 1" else
                       "### " + t if st == "Heading 2" else
                       "- " + t if st.startswith("List") else t)
        elif child.tag == qn('w:tbl'):
            tb = tbl_map.get(child)
            if tb is None:
                continue
            for row in tb.rows:
                out.append("| " + " | ".join(c.text.strip().replace("\n", " ") for c in row.cells) + " |")
    return "\n".join(out)

def main():
    items = []
    for f in sorted(glob.glob(os.path.join(C.DOCDIR, "**", "*.docx"), recursive=True)):
        b = os.path.basename(f)
        if b in EXCLUDE or b.startswith("~$"):
            continue
        rel = os.path.relpath(f, C.DOCDIR).replace("\\", "/")
        d = Document(f)
        title = next((p.text.strip() for p in d.paragraphs
                      if p.style and p.style.name == "Title" and p.text.strip()),
                     os.path.splitext(b)[0])
        items.append({"source": title, "file": rel, "text": doc_to_text(f)})
    payload = json.dumps(items, ensure_ascii=False)
    b64 = base64.b64encode(payload.encode("utf-8")).decode("ascii")
    open(C.B64, "w", encoding="ascii").write(b64)
    print(f"Documentos: {len(items)} | chars: {sum(len(i['text']) for i in items)} | base64 len: {len(b64)}")
    print("Guardado:", C.B64)

if __name__ == "__main__":
    main()
